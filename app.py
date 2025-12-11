import streamlit as st
import pandas as pd
import math
from core.models import LoadInput, InstallationParams, ConduitType, ConductorMaterial, InsulationRating
from core.converters import convert_power_unit, convert_length_unit
from standards.nec_logic import NECLogic
from standards.nec_tables import NEC_310_16_COPPER
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# --- Page Config ---
st.set_page_config(
    page_title="Calculadora de Tableros (NEC)",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- Custom CSS ---
st.markdown("""
<style>
    .reportview-container { background: #f0f2f6; }
    .main-header { font-family: 'Inter', sans-serif; color: #1E3A8A; font-weight: 700; }
    .metric-card { background-color: white; padding: 20px; border-radius: 10px; box-shadow: 0 4px 6px rgba(0,0,0,0.1); text-align: center; }
    .stButton>button { width: 100%; border-radius: 8px; font-weight: 600; }
    .stDataFrame { border-radius: 10px; overflow: hidden; }
</style>
""", unsafe_allow_html=True)

# --- Session State Init ---
if 'loads' not in st.session_state:
    st.session_state.loads = []

if "voltage_input" not in st.session_state:
    st.session_state.voltage_input = 480.0

def on_phase_change():
    if st.session_state.phases_input == 1:
        st.session_state.voltage_input = 220.0
    else:
        st.session_state.voltage_input = 480.0

# --- Sidebar: Global Settings ---
with st.sidebar:
    st.image("https://cdn-icons-png.flaticon.com/512/2933/2933886.png", width=60)
    st.title("Configuración Global")
    st.info("💡 Pase el mouse sobre el icono (?) junto a los campos para ver referencias a las tablas NEC.")
    
    st.markdown("---")
    st.subheader("📥 Importación Masiva")
    
    # 1. Download Template
    def get_template():
        # Create a sample DataFrame
        data = {
            "Nombre": ["Motor Bomba", "Iluminación Hall"],
            "Cantidad": [1, 15],
            "Potencia": [10, 100],
            "UnidadPotencia": ["HP", "W"],
            "Voltaje": [480, 220],
            "Fases": [3, 1],
            "FP": [0.85, 0.95],
            "Longitud": [50, 20],
            "UnidadLongitud": ["m", "m"],
            "EsMotor": ["SI", "NO"],
            "EsContinuo": ["NO", "SI"],
            "TempAmb": [30, 30],
            "TipoDucto": ["ACERO", "PVC"],
            "Agrupamiento": [3, 3]
        }
        df = pd.DataFrame(data)
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Plantilla')
        return output.getvalue()
        
    st.download_button(
        "📄 Descargar Plantilla Excel",
        data=get_template(),
        file_name="plantilla_cargas_nec.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        help="Descargue esta plantilla, llénela con sus cargas y súbala abajo."
    )
    
    # 2. Upload File
    uploaded_file = st.file_uploader("Subir Excel", type=["xlsx"], help="Asegúrese de respetar los nombres de columnas de la plantilla.")
    
    if uploaded_file:
        if st.button("Procesar Archivo"):
            try:
                df_in = pd.read_excel(uploaded_file)
                success_count = 0
                errors = []
                
                for idx, row in df_in.iterrows():
                    try:
                        # Parsing
                        p_unit = str(row.get("UnidadPotencia", "W")).strip()
                        l_unit = str(row.get("UnidadLongitud", "m")).strip()
                        watts, amps_override = convert_power_unit(float(row["Potencia"]), p_unit, float(row["Voltaje"]), int(row["Fases"]), float(row["FP"]))
                        len_m = convert_length_unit(float(row["Longitud"]), l_unit)
                        
                        # Booleans
                        is_mot = str(row.get("EsMotor", "NO")).strip().upper() == "SI"
                        is_cont = str(row.get("EsContinuo", "NO")).strip().upper() == "SI"
                        
                        # Conduit
                        c_str = str(row.get("TipoDucto", "PVC")).strip().upper()
                        c_type = ConduitType.STEEL if "ACERO" in c_str or "STEEL" in c_str else (ConduitType.ALUMINUM if "ALUM" in c_str else ConduitType.PVC)
                        
                        load_obj = LoadInput(
                            name=str(row["Nombre"]), 
                            power_watts=watts, 
                            voltage=float(row["Voltaje"]), 
                            phases=int(row["Fases"]),
                            is_continuous=is_cont, 
                            is_motor=is_mot, 
                            power_factor=float(row["FP"]), 
                            quantity=int(row.get("Cantidad", 1)), 
                            override_amps=amps_override
                        )
                        
                        inst_params = InstallationParams(
                            length_meters=len_m,
                            conduit_type=c_type,
                            ambient_temp_c=float(row.get("TempAmb", 30)),
                            raceway_count=int(row.get("Agrupamiento", 3)),
                            conductor_material=ConductorMaterial.COPPER,
                            insulation_rating=InsulationRating.TEMP_75 # Default import as 75 for safety, or add col
                        )
                        
                        st.session_state.loads.append({"load": load_obj, "params": inst_params})
                        success_count += 1
                        
                    except Exception as e:
                        errors.append(f"Fila {idx+2}: {str(e)}")
                
                if success_count > 0:
                    st.success(f"✅ {success_count} cargas importadas.")
                    if errors:
                        st.warning(f"⚠️ {len(errors)} errores: {'; '.join(errors)}")
                    st.rerun()
                else:
                    st.error("No se pudieron importar cargas. Revise el formato.")
                    
            except Exception as e:
                st.error(f"Error leyendo archivo: {e}")

# --- Main Area ---
st.markdown("<h1 class='main-header'>⚡ Calculadora de Circuitos (NEC Edition v2.2)</h1>", unsafe_allow_html=True)
st.markdown("---")

col1, col2 = st.columns([1, 2])

# --- Section 1: Add New Load ---
with col1:
    with st.container():
        st.subheader("➕ Agregar Carga Manual")
        
        name = st.text_input("Nombre de Carga", placeholder="Ej. Motor Bomba 1")
        
        c1, c2 = st.columns(2)
        with c1:
            qty = st.number_input("Cantidad", min_value=1, value=1, step=1, help="Multiplicador para cargas idénticas.")
        with c2:
            phases = st.radio("Fases", [1, 3], horizontal=True, key="phases_input", on_change=on_phase_change)
        
        c3, c4 = st.columns([2, 1])
        with c3:
            power_val = st.number_input("Potencia", min_value=0.0, value=0.0, step=0.1)
        with c4:
            power_unit = st.selectbox("Unidad", ["W", "KW", "HP", "A", "KVA"], help="HP: Se convierte a Watts (x746). KVA: Se aplica Factor de Potencia.")
            
        voltage = st.number_input("Voltaje (V)", step=10.0, key="voltage_input")
        
        pf = st.number_input("Factor de Potencia", value=0.9, min_value=0.1, max_value=1.0, step=0.05, help="Afecta el cálculo de corriente en AC y la impedancia efectiva (Z) para caída de tensión.")
        
        st.markdown("**Detalles del Circuito Derivado**")
        length_val = st.number_input("Longitud", value=10.0, help="Distancia unidireccional. Usado para Caída de Tensión.")
        length_unit = st.selectbox("Unidad Long.", ["m", "ft"], index=0)
        
        with st.expander("Opciones Avanzadas (NEC) - Ambiente y Ducto", expanded=True):
            is_motor = st.checkbox("Es Motor", value=False, help="Aplica reglas de NEC 430.22 (125% FLA).")
            is_continuous = st.checkbox("Carga Continua (>3h)", value=False, help="NEC 215.2(A)(1): Requiere 125% de capacidad.")
            
            st.markdown("---")
            st.markdown("**Condiciones de Instalación**")
            
            temp_c = st.number_input("Temperatura Ambiente (°C)", value=30.0, step=1.0, min_value=-50.0, max_value=120.0, 
                                     help="NEC Tabla 310.15(B)(1): Corrección por temperatura.")
            
            conduit_options = {
                "PVC (No Magnético)": ConduitType.PVC,
                "Acero (Magnético)": ConduitType.STEEL,
                "Aluminio": ConduitType.ALUMINUM
            }
            conduit_label = st.selectbox("Tipo de Ducto", list(conduit_options.keys()), 
                                         help="NEC Table 9: Material afecta reactancia (X).")
            local_conduit = conduit_options[conduit_label]
            
            temp_rating_choice = st.radio("Temp. Aislamiento", ["75°C (THWN)", "90°C (THHN/THWN-2)"], 
                                          help="NEC 310.16 / 110.14(C): Ampacidad base.")
            rating_enum = InsulationRating.TEMP_90 if "90" in temp_rating_choice else InsulationRating.TEMP_75
            
            grouping_count = st.number_input("Total Conductores en Ducto", min_value=1, value=3, 
                                             help="NEC Tabla 310.15(C)(1): Factor de agrupamiento.")
        
        if st.button("Agregar Carga", type="primary"):
            if name:
                watts, amps_override = convert_power_unit(power_val, power_unit, voltage, phases, pf)
                length_m = convert_length_unit(length_val, length_unit)
                
                load_obj = LoadInput(
                    name=name, power_watts=watts, voltage=voltage, phases=phases,
                    is_continuous=is_continuous, is_motor=is_motor, power_factor=pf,
                    quantity=qty, override_amps=amps_override
                )
                
                inst_params = InstallationParams(
                    length_meters=length_m,
                    conduit_type=local_conduit,
                    ambient_temp_c=temp_c,
                    raceway_count=grouping_count,
                    conductor_material=ConductorMaterial.COPPER,
                    insulation_rating=rating_enum
                )
                
                st.session_state.loads.append({"load": load_obj, "params": inst_params})
                st.success(f"Carga '{name}' agregada.")
            else:
                st.error("El nombre es obligatorio.")

# --- Section 2: Results Table ---
with col2:
    st.subheader("📋 Tabla de Cargas y Resultados")
    
    if not st.session_state.loads:
        st.info("No hay cargas agregadas. Utilice el formulario de la izquierda o importe desde Excel en el menú lateral.")
    else:
        results_data = []
        feeder_loads = []
        
        for i, item in enumerate(st.session_state.loads):
            l = item["load"]
            p = item["params"]
            
            res = NECLogic.select_conductor_and_breaker(l, p)
            
            if l.override_amps:
                disp_amps = l.override_amps
            else:
                 if l.phases == 1:
                    disp_amps = l.power_watts / (l.voltage * l.power_factor)
                 else:
                    disp_amps = l.power_watts / (math.sqrt(3) * l.voltage * l.power_factor)
            
            c_disp = "PVC"
            if p.conduit_type == ConduitType.STEEL: c_disp = "Acero"
            elif p.conduit_type == ConduitType.ALUMINUM: c_disp = "Alum"

            results_data.append({
                "Cant": l.quantity,
                "Nombre": l.name,
                "Amps": round(disp_amps, 1),
                "Calibre": res.size,
                "Ampacidad": res.ampacity,
                "Breaker": res.breaker_rating,
                "% VD": float(f"{res.voltage_drop_percent:.2f}"),
                "T.Amb": f"{p.ambient_temp_c}°C",
                "Ducto": c_disp,
                "Rating": f"{p.insulation_rating.value}C"
            })
            
            feeder_loads.append(l)

        df = pd.DataFrame(results_data)
        
        st.dataframe(
            df.style.map(lambda x: 'color: red; font-weight: bold' if isinstance(x, float) and x > 3.0 else '', subset=['% VD']),
            use_container_width=True,
            hide_index=True
        )
        
        if st.button("🗑️ Limpiar Todo", type="secondary"):
            st.session_state.loads = []
            st.rerun()

# --- Section 3: Main Feeder ---
st.markdown("---")
st.subheader("🏢 Alimentador Principal (Main Feeder)")

if st.session_state.loads:
    feeder_res = NECLogic.calculate_main_feeder(feeder_loads)
    
    m1, m2, m3 = st.columns(3)
    with m1:
        st.metric("Corriente Total Est.", f"{feeder_res['total_amps']:.1f} A", help="Suma de cargas según NEC 430.24 (Motores) y 215.2 (F. Demanda).")
    with m2:
        st.metric("Conductor Sugerido", feeder_res['cable_size'], help="Seleccionado de NEC Tabla 310.16 (75°C).")
    with m3:
        runs = feeder_res['parallel_runs']
        desc = feeder_res['description']
        st.metric("Configuración", "Simple" if runs==1 else f"{runs}x Paralelo", delta=desc, delta_color="off", help="Configuración paralela si Corriente > 400A (NEC 310.10(H)).")
    
    st.info(f"**Nota:** Cálculo basado en NEC 430.24 (Motores) y 215.2 (Continuo + Demand). {feeder_res['description']}")

# --- Section 4: Export ---
def generate_excel():
    output = io.BytesIO()
    wb = Workbook()
    
    # Sheet 1: Branch Circuits
    ws = wb.active
    ws.title = "Circuitos Derivados"
    ws.append(["Cant", "Nombre", "Potencia W", "Voltaje", "Fases", "Amps", "Calibre", "Ampacidad", "Breaker", "% VD", "T.Amb", "Tipo Ducto", "Insul.", "Notas"])
    
    for i, item in enumerate(st.session_state.loads):
        l = item["load"]
        p = item["params"]
        r = NECLogic.select_conductor_and_breaker(l, p)
        
        if l.override_amps:
            a = l.override_amps
        else:
             a = l.power_watts / (l.voltage * (1.732 if l.phases==3 else 1) * l.power_factor)
             
        ws.append([l.quantity, l.name, l.power_watts, l.voltage, l.phases, round(a,2), r.size, r.ampacity, r.breaker_rating, r.voltage_drop_percent, p.ambient_temp_c, p.conduit_type.value, p.insulation_rating.value, r.reference_notes])

    # Sheet 2: Feeder
    if st.session_state.loads:
        ws2 = wb.create_sheet("Alimentador")
        fres = NECLogic.calculate_main_feeder([item["load"] for item in st.session_state.loads])
        ws2.append(["Parametro", "Valor"])
        ws2.append(["Total Amps", fres['total_amps']])
        ws2.append(["Cable", fres['cable_size']])
        ws2.append(["Config", fres['description']])

    wb.save(output)
    return output.getvalue()

if st.session_state.loads:
    excel_data = generate_excel()
    st.download_button(
        label="📥 Descargar Reporte Excel",
        data=excel_data,
        file_name="memoria_calculo_nec_v2.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
